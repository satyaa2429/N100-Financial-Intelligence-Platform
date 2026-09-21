from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Font

from src.screener.engine import ScreenerEngine
from src.screener.ranking import build_ranking


OUTPUT = Path("output/screener_output.xlsx")


def format_workbook(path: Path) -> None:
    """Apply presentation formatting to the screener workbook."""
    workbook = load_workbook(path)

    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions

        for cell in sheet[1]:
            cell.font = Font(bold=True)

        for column_cells in sheet.columns:
            width = max(
                len(str(cell.value)) if cell.value is not None else 0
                for cell in column_cells
            )
            letter = column_cells[0].column_letter
            sheet.column_dimensions[letter].width = min(width + 2, 35)

        headers = {
            cell.value: cell.column
            for cell in sheet[1]
        }

        for score_column in (
            "profitability_score",
            "growth_score",
            "valuation_score",
            "composite_score",
        ):
            if score_column in headers and sheet.max_row > 1:
                col = headers[score_column]
                letter = sheet.cell(1, col).column_letter

                sheet.conditional_formatting.add(
                    f"{letter}2:{letter}{sheet.max_row}",
                    ColorScaleRule(
                        start_type="min",
                        start_color="F8696B",
                        mid_type="percentile",
                        mid_value=50,
                        mid_color="FFEB84",
                        end_type="max",
                        end_color="63BE7B",
                    ),
                )

    workbook.save(path)


def main() -> None:
    """Generate the formatted screener workbook."""
    ranking = build_ranking()
    engine = ScreenerEngine()

    preset_names = [
        "quality",
        "value",
        "growth",
        "dividend",
        "momentum",
        "debt_free",
    ]

    sheet_names = {
        "quality": "Quality",
        "value": "Value",
        "growth": "Growth",
        "dividend": "Dividend",
        "momentum": "Momentum",
        "debt_free": "Debt-Free",
    }

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)

    with pd.ExcelWriter(
        OUTPUT,
        engine="openpyxl",
    ) as writer:

        ranking.to_excel(
            writer,
            sheet_name="Composite Ranking",
            index=False,
        )

        for preset in preset_names:
            preset_result = engine.run_preset(preset)

            ids = set(preset_result["company_id"])

            ranked_preset = (
                ranking[
                    ranking["company_id"].isin(ids)
                ]
                .sort_values(
                    ["composite_score", "company_id"],
                    ascending=[False, True],
                )
                .reset_index(drop=True)
            )

            ranked_preset.to_excel(
                writer,
                sheet_name=sheet_names[preset],
                index=False,
            )

        for sector, group in ranking.groupby("broad_sector"):
            group.sort_values(
                "sector_rank"
            ).to_excel(
                writer,
                sheet_name=str(sector)[:31],
                index=False,
            )

    format_workbook(OUTPUT)

    print("Workbook:", OUTPUT.resolve())
    print("Composite rows:", len(ranking))

    for preset in preset_names:
        print(
            sheet_names[preset],
            "=>",
            len(engine.run_preset(preset)),
            "rows",
        )


if __name__ == "__main__":
    main()
