import os
import sqlite3
import pandas as pd

from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter


DB_PATH = r"data\nifty100.db"
CAPITAL_FILE = r"output\capital_allocation.csv"

OUTPUT_XLSX = r"output\cashflow_intelligence.xlsx"
DISTRESS_CSV = r"output\distress_alerts.csv"

os.makedirs("output", exist_ok=True)


# ---------------------------------------------------------
# Helpers
# ---------------------------------------------------------

def quality_label(value):
    """Return the cash-flow quality label for a numeric score."""
    if pd.isna(value):
        return "Unavailable"

    if value > 1.0:
        return "High Quality Earnings"

    if value < 0.5:
        return "Accrual Risk"

    return "Moderate"


def calculate_fcf_cagr(values, years):
    """
    FCF CAGR using annual observations.

    Positive start/end -> normal CAGR.
    Negative/non-positive start to positive end -> turnaround.
    Otherwise CAGR unavailable.
    """

    series = pd.Series(values).dropna()

    required = years + 1

    if len(series) < required:
        return None, False

    start = float(series.iloc[-required])
    end = float(series.iloc[-1])

    if start > 0 and end > 0:
        value = (
            (end / start) ** (1 / years)
            - 1
        ) * 100

        return value, False

    if start <= 0 < end:
        return None, True

    return None, False


# ---------------------------------------------------------
# Load existing capital-allocation output
# ---------------------------------------------------------

capital = pd.read_csv(CAPITAL_FILE)

capital["year_num"] = pd.to_numeric(
    capital["year"].astype(str).str[:4],
    errors="coerce",
)

capital = capital.sort_values(
    ["company_id", "year_num"]
)


# ---------------------------------------------------------
# Load company, sector and debt information
# ---------------------------------------------------------

con = sqlite3.connect(DB_PATH)

companies = pd.read_sql_query(
    """
    SELECT
        id AS company_id,
        company_name
    FROM companies
    """,
    con,
)

sectors = pd.read_sql_query(
    """
    SELECT
        company_id,
        broad_sector
    FROM sectors
    """,
    con,
)

debt = pd.read_sql_query(
    """
    SELECT
        company_id,
        year,
        total_debt_cr
    FROM financial_ratios
    WHERE UPPER(year) <> 'TTM'
    """,
    con,
)

con.close()

debt["year_num"] = pd.to_numeric(
    debt["year"].astype(str).str[:4],
    errors="coerce",
)

debt = debt.sort_values(
    ["company_id", "year_num"]
)


# ---------------------------------------------------------
# Debt latest / previous year
# ---------------------------------------------------------

debt_rows = []

for company_id, group in debt.groupby("company_id"):

    valid = (
        group.dropna(subset=["total_debt_cr"])
        .sort_values("year_num")
    )

    if valid.empty:
        latest_debt = None
        previous_debt = None

    else:
        latest_debt = valid.iloc[-1]["total_debt_cr"]

        previous_debt = (
            valid.iloc[-2]["total_debt_cr"]
            if len(valid) >= 2
            else None
        )

    debt_rows.append(
        {
            "company_id": company_id,
            "latest_total_debt_cr": latest_debt,
            "previous_total_debt_cr": previous_debt,
        }
    )

debt_summary = pd.DataFrame(debt_rows)


# ---------------------------------------------------------
# Build company cash-flow intelligence
# ---------------------------------------------------------

summary_rows = []

for company_id, history in capital.groupby("company_id"):

    history = history.sort_values("year_num")

    latest = history.iloc[-1]

    # ----- CFO Quality: five-year average -----

    cfo_pat_5yr = (
        history["cfo_pat_ratio"]
        .dropna()
        .tail(5)
    )

    if len(cfo_pat_5yr) > 0:
        cfo_quality_score = cfo_pat_5yr.mean()
    else:
        cfo_quality_score = None

    cfo_quality = quality_label(
        cfo_quality_score
    )

    # ----- FCF CAGR -----

    fcf_5yr, fcf_5yr_turnaround = (
        calculate_fcf_cagr(
            history["free_cash_flow_cr"],
            5,
        )
    )

    fcf_10yr, fcf_10yr_turnaround = (
        calculate_fcf_cagr(
            history["free_cash_flow_cr"],
            10,
        )
    )

    summary_rows.append(
        {
            "company_id": company_id,
            "latest_year": latest["year"],

            "free_cash_flow_cr":
                latest["free_cash_flow_cr"],

            "cfo_quality_score_5yr_avg":
                cfo_quality_score,

            "cfo_quality_label":
                cfo_quality,

            "fcf_cagr_5yr_pct":
                fcf_5yr,

            "fcf_cagr_5yr_turnaround":
                fcf_5yr_turnaround,

            "fcf_cagr_10yr_pct":
                fcf_10yr,

            "fcf_cagr_10yr_turnaround":
                fcf_10yr_turnaround,

            "capex_intensity_pct":
                latest["capex_intensity_pct"],

            "capex_intensity_label":
                latest["capex_intensity_label"],

            "fcf_conversion_pct":
                latest["fcf_conversion_pct"],

            "fcf_conversion_label":
                latest["fcf_conversion_label"],

            "operating_activity_cr":
                latest["operating_activity"],

            "investing_activity_cr":
                latest["investing_activity"],

            "cfo_sign":
                latest["cfo_sign"],

            "cfi_sign":
                latest["cfi_sign"],

            "cff_sign":
                latest["cff_sign"],

            "capital_allocation_label":
                latest["pattern_label"],

            "distress_flag":
                bool(latest["distress_flag"]),
        }
    )


summary = pd.DataFrame(summary_rows)

summary = (
    summary
    .merge(
        companies,
        on="company_id",
        how="left",
    )
    .merge(
        sectors,
        on="company_id",
        how="left",
    )
    .merge(
        debt_summary,
        on="company_id",
        how="left",
    )
)


# ---------------------------------------------------------
# Deleveraging flag
# CFF < 0 AND debt declining YoY
# ---------------------------------------------------------

summary["deleveraging_flag"] = (
    (summary["cff_sign"] == "-")
    &
    summary["latest_total_debt_cr"].notna()
    &
    summary["previous_total_debt_cr"].notna()
    &
    (
        summary["latest_total_debt_cr"]
        <
        summary["previous_total_debt_cr"]
    )
)


# ---------------------------------------------------------
# Column order
# ---------------------------------------------------------

summary = summary[
    [
        "company_id",
        "company_name",
        "broad_sector",
        "latest_year",

        "operating_activity_cr",
        "investing_activity_cr",
        "free_cash_flow_cr",

        "cfo_quality_score_5yr_avg",
        "cfo_quality_label",

        "fcf_cagr_5yr_pct",
        "fcf_cagr_5yr_turnaround",

        "fcf_cagr_10yr_pct",
        "fcf_cagr_10yr_turnaround",

        "capex_intensity_pct",
        "capex_intensity_label",

        "fcf_conversion_pct",
        "fcf_conversion_label",

        "latest_total_debt_cr",
        "previous_total_debt_cr",
        "deleveraging_flag",

        "cfo_sign",
        "cfi_sign",
        "cff_sign",

        "capital_allocation_label",
        "distress_flag",
    ]
].sort_values("company_id")


# ---------------------------------------------------------
# Distress alerts
# ---------------------------------------------------------

distress = summary[
    summary["distress_flag"] == True
].copy()

distress.to_csv(
    DISTRESS_CSV,
    index=False,
)


# ---------------------------------------------------------
# Capital-allocation matrix
# ---------------------------------------------------------

pattern_matrix = (
    summary["capital_allocation_label"]
    .value_counts(dropna=False)
    .rename_axis("capital_allocation_label")
    .reset_index(name="company_count")
)

pattern_matrix["percentage_of_universe"] = (
    pattern_matrix["company_count"]
    / len(summary)
    * 100
)


# ---------------------------------------------------------
# Export workbook
# ---------------------------------------------------------

with pd.ExcelWriter(
    OUTPUT_XLSX,
    engine="openpyxl",
) as writer:

    summary.to_excel(
        writer,
        sheet_name="Cashflow Intelligence",
        index=False,
    )

    pattern_matrix.to_excel(
        writer,
        sheet_name="Pattern Matrix",
        index=False,
    )

    distress.to_excel(
        writer,
        sheet_name="Distress Alerts",
        index=False,
    )

    capital.drop(
        columns=["year_num"],
        errors="ignore",
    ).to_excel(
        writer,
        sheet_name="Cashflow History",
        index=False,
    )


# ---------------------------------------------------------
# Basic professional formatting
# ---------------------------------------------------------

wb = load_workbook(OUTPUT_XLSX)

header_fill = PatternFill(
    "solid",
    fgColor="1F4E78",
)

header_font = Font(
    color="FFFFFF",
    bold=True,
)

for ws in wb.worksheets:

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
        )

    for column in ws.columns:

        column_letter = get_column_letter(
            column[0].column
        )

        max_length = 0

        for cell in column[:200]:

            value = "" if cell.value is None else str(cell.value)

            max_length = max(
                max_length,
                len(value),
            )

        ws.column_dimensions[
            column_letter
        ].width = min(
            max(max_length + 2, 12),
            32,
        )


wb.save(OUTPUT_XLSX)


# ---------------------------------------------------------
# Validation summary
# ---------------------------------------------------------

print("Created:", OUTPUT_XLSX)
print("Created:", DISTRESS_CSV)

print("\nCompanies covered:", summary["company_id"].nunique())
print("Summary rows:", len(summary))

print(
    "Distress companies:",
    summary["distress_flag"].sum(),
)

print(
    "Deleveraging companies:",
    summary["deleveraging_flag"].sum(),
)

print(
    "Capital-allocation patterns:",
    summary["capital_allocation_label"].nunique(),
)

print("\nCFO quality distribution:")
print(
    summary["cfo_quality_label"]
    .value_counts()
    .to_string()
)

print("\nPattern distribution:")
print(
    summary["capital_allocation_label"]
    .value_counts()
    .to_string()
)