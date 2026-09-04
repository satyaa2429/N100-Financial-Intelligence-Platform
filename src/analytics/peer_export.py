"""Export Sprint 3 Day 20 peer comparison workbook."""

from pathlib import Path
import sqlite3
import re

import pandas as pd

from openpyxl import Workbook
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


DB_PATH = Path("data/nifty100.db")
OUTPUT_PATH = Path("output/peer_comparison.xlsx")
YEAR = "2024-03"

METRICS = [
    "net_profit_margin_pct",
    "operating_profit_margin_pct",
    "return_on_equity_pct",
    "return_on_capital_employed_pct",
    "debt_to_equity",
    "interest_coverage",
    "asset_turnover",
    "free_cash_flow_cr",
    "cash_from_operations_cr",
    "earnings_per_share",
    "book_value_per_share",
    "dividend_payout_ratio_pct",
    "revenue_cagr_5yr",
    "pat_cagr_5yr",
    "eps_cagr_5yr",
    "cfo_pat_ratio",
    "capex_intensity_pct",
    "fcf_conversion_pct",
    "pe_ratio",
    "pb_ratio",
]


def load_data():
    with sqlite3.connect(DB_PATH) as conn:

        percentiles = pd.read_sql_query(
            """
            SELECT
                company_id,
                peer_group,
                metric,
                value,
                percentile_rank
            FROM peer_percentiles
            WHERE year = ?
            """,
            conn,
            params=(YEAR,),
        )

        companies = pd.read_sql_query(
            """
            SELECT
                id AS company_id,
                company_name
            FROM companies
            """,
            conn,
        )

        peers = pd.read_sql_query(
            """
            SELECT
                peer_group_name AS peer_group,
                company_id,
                is_benchmark
            FROM peer_groups
            """,
            conn,
        )

    return percentiles, companies, peers


def safe_sheet_name(name):
    name = re.sub(r'[:\\/?*\[\]]', "_", name)
    return name[:31]


def build_group_table(
    group,
    percentiles,
    companies,
    peers,
):
    members = peers[
        peers["peer_group"] == group
    ].copy()

    members = members.merge(
        companies,
        on="company_id",
        how="left",
    )

    group_data = percentiles[
        percentiles["peer_group"] == group
    ].copy()

    values = group_data.pivot(
        index="company_id",
        columns="metric",
        values="value",
    )

    ranks = group_data.pivot(
        index="company_id",
        columns="metric",
        values="percentile_rank",
    )

    output = members[
        [
            "company_id",
            "company_name",
            "is_benchmark",
        ]
    ].copy()

    output["is_benchmark"] = output[
        "is_benchmark"
    ].map(
        {
            1: "Yes",
            0: "No",
        }
    )

    for metric in METRICS:

        value_col = f"{metric} | Value"
        pct_col = f"{metric} | Percentile"

        output[value_col] = output[
            "company_id"
        ].map(
            values[metric]
            if metric in values.columns
            else pd.Series(dtype=float)
        )

        output[pct_col] = output[
            "company_id"
        ].map(
            ranks[metric]
            if metric in ranks.columns
            else pd.Series(dtype=float)
        )

    return output


def style_sheet(ws, row_count, col_count):

    header_fill = PatternFill(
        "solid",
        fgColor="1F4E78",
    )

    benchmark_fill = PatternFill(
        "solid",
        fgColor="D9EAF7",
    )

    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = Font(
            color="FFFFFF",
            bold=True,
        )
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )

    ws.row_dimensions[1].height = 42

    ws.freeze_panes = "D2"
    ws.auto_filter.ref = (
        f"A1:{get_column_letter(col_count)}{row_count}"
    )

    for row in range(2, row_count + 1):

        if ws.cell(row, 3).value == "Yes":
            for col in range(
                1,
                col_count + 1,
            ):
                ws.cell(
                    row,
                    col,
                ).fill = benchmark_fill

            ws.cell(
                row,
                1,
            ).font = Font(bold=True)

    for col in range(
        4,
        col_count + 1,
    ):

        header = str(
            ws.cell(1, col).value
        )

        if header.endswith(
            "| Percentile"
        ):
            ws.cell(
                1,
                col,
            ).fill = PatternFill(
                "solid",
                fgColor="548235",
            )

            for row in range(
                2,
                row_count + 1,
            ):
                ws.cell(
                    row,
                    col,
                ).number_format = "0.0%"

            rng = (
                f"{get_column_letter(col)}2:"
                f"{get_column_letter(col)}{row_count}"
            )

            ws.conditional_formatting.add(
                rng,
                ColorScaleRule(
                    start_type="num",
                    start_value=0,
                    start_color="F8696B",
                    mid_type="num",
                    mid_value=0.5,
                    mid_color="FFEB84",
                    end_type="num",
                    end_value=1,
                    end_color="63BE7B",
                ),
            )

        else:
            for row in range(
                2,
                row_count + 1,
            ):
                cell = ws.cell(
                    row,
                    col,
                )

                if isinstance(
                    cell.value,
                    (int, float),
                ):
                    cell.number_format = "0.00"

    ws.column_dimensions["A"].width = 16
    ws.column_dimensions["B"].width = 34
    ws.column_dimensions["C"].width = 13

    for col in range(
        4,
        col_count + 1,
    ):
        ws.column_dimensions[
            get_column_letter(col)
        ].width = 21

    for row in ws.iter_rows(
        min_row=2,
        max_row=row_count,
    ):
        for cell in row:
            cell.alignment = Alignment(
                vertical="center",
            )


def main():

    percentiles, companies, peers = (
        load_data()
    )

    groups = sorted(
        peers["peer_group"].unique()
    )

    wb = Workbook()

    default = wb.active
    wb.remove(default)

    summary = []

    for group in groups:

        table = build_group_table(
            group,
            percentiles,
            companies,
            peers,
        )

        ws = wb.create_sheet(
            safe_sheet_name(group)
        )

        ws.append(
            list(table.columns)
        )

        for row in table.itertuples(
            index=False,
            name=None,
        ):
            ws.append(
                [
                    None
                    if pd.isna(value)
                    else value
                    for value in row
                ]
            )

        style_sheet(
            ws,
            ws.max_row,
            ws.max_column,
        )

        benchmark_count = int(
            (
                table["is_benchmark"]
                == "Yes"
            ).sum()
        )

        summary.append(
            (
                group,
                len(table),
                benchmark_count,
            )
        )

    OUTPUT_PATH.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    wb.save(OUTPUT_PATH)

    print("Workbook:", OUTPUT_PATH.resolve())
    print("Sheets:", len(wb.sheetnames))
    print("Sheet names:", wb.sheetnames)

    total_companies = 0

    for group, members, benchmarks in summary:
        print(
            group,
            "=> companies:",
            members,
            "| benchmark:",
            benchmarks,
        )
        total_companies += members

    print("Total peer memberships:", total_companies)
    print(
        "Metrics per sheet:",
        len(METRICS),
    )


if __name__ == "__main__":
    main()
